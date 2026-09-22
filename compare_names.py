#!/usr/bin/env python3
from __future__ import annotations

import argparse
import difflib
import re
import shutil
import unicodedata
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Set, Tuple

from openpyxl import load_workbook


def normalize_name(value: object) -> str:
    """Normalize names for case-insensitive and accent-insensitive matching."""
    if value is None:
        return ""
    text = str(value).strip()
    if not text:
        return ""

    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.lower()
    text = re.sub(r"[^a-z0-9\s'-]", "", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def looks_like_first_name_header(header: str) -> bool:
    return bool(re.search(r"\b(first|given)\b", header))


def looks_like_last_name_header(header: str) -> bool:
    return bool(re.search(r"\b(last|family|surname)\b", header))


def detect_name_columns(header_cells: Iterable[object]) -> Tuple[Optional[int], Optional[int], Dict[int, str]]:
    headers: Dict[int, str] = {}
    first_col = None
    last_col = None

    for col_index, value in enumerate(header_cells, start=1):
        if value is None:
            continue

        raw = str(value).strip()
        if not raw:
            continue

        normalized = normalize_name(raw)
        headers[col_index] = raw

        if first_col is None and looks_like_first_name_header(normalized):
            first_col = col_index
        if last_col is None and looks_like_last_name_header(normalized):
            last_col = col_index

    return first_col, last_col, headers


def resolve_input_path(base_dir: Path, explicit: Optional[str], pattern: str) -> Path:
    if explicit:
        path = Path(explicit)
        if not path.is_absolute():
            path = (base_dir / path).resolve()
        if not path.exists():
            raise FileNotFoundError(f"Input file not found: {path}")
        return path

    candidates = sorted(base_dir.glob(pattern))
    if not candidates:
        raise FileNotFoundError(f"No file matched pattern: {pattern}")
    if len(candidates) > 1:
        raise RuntimeError(
            "Multiple files matched pattern. Please pass an explicit path. "
            f"Pattern: {pattern}; Matches: {', '.join(str(p) for p in candidates)}"
        )
    return candidates[0].resolve()


def tokenize_name(name: str) -> List[str]:
    return [part for part in re.split(r"[\s'-]+", name) if part]


def has_similar_token(source_name: str, target_name: str, threshold: float = 0.84) -> bool:
    """Return True if any meaningful token in source_name closely matches a token in target_name."""
    source_tokens = [tok for tok in tokenize_name(source_name) if len(tok) >= 3]
    target_tokens = [tok for tok in tokenize_name(target_name) if len(tok) >= 3]

    if not source_tokens or not target_tokens:
        return False

    for source_tok in source_tokens:
        for target_tok in target_tokens:
            if source_tok == target_tok:
                return True
            if difflib.SequenceMatcher(a=source_tok, b=target_tok).ratio() >= threshold:
                return True
    return False


def is_similar_name(a: str, b: str) -> bool:
    """Return True when two normalized names are close enough for manual review."""
    if not a or not b:
        return False
    if a == b:
        return True

    # Fast path for substring-style differences like "cho kwong charlie" vs "charlie".
    if a in b or b in a:
        return True

    tokens_a = tokenize_name(a)
    tokens_b = tokenize_name(b)
    if not tokens_a or not tokens_b:
        return False

    overlap = set(tokens_a) & set(tokens_b)
    min_len = min(len(set(tokens_a)), len(set(tokens_b)))
    if min_len > 0 and len(overlap) >= min_len:
        return True

    return difflib.SequenceMatcher(a=a, b=b).ratio() >= 0.72


def display_full_name(first_display: str, last_display: str) -> str:
    return " ".join(part for part in [first_display.strip(), last_display.strip()] if part).strip()


def collect_rsvp_people(rsvp_path: Path, first_col: int = 7, last_col: int = 8) -> List[Dict[str, str]]:
    wb = load_workbook(rsvp_path, data_only=True)
    ws = wb.active

    people: List[Dict[str, str]] = []
    for row in range(2, ws.max_row + 1):
        first_raw = ws.cell(row=row, column=first_col).value
        last_raw = ws.cell(row=row, column=last_col).value
        first = normalize_name(first_raw)
        last = normalize_name(last_raw)
        first_display = str(first_raw).strip() if first_raw is not None else ""
        last_display = str(last_raw).strip() if last_raw is not None else ""
        full_display = display_full_name(first_display, last_display)

        if not first and not last:
            continue

        people.append(
            {
                "first": first,
                "last": last,
                "full_display": full_display or (first_display or last_display),
            }
        )

    return people


def build_exact_name_sets(rsvp_people: List[Dict[str, str]]) -> Tuple[Set[str], Set[str]]:
    first_names: Set[str] = set()
    last_names: Set[str] = set()

    for person in rsvp_people:
        if person["first"]:
            first_names.add(person["first"])
        if person["last"]:
            last_names.add(person["last"])

    return first_names, last_names


def find_maybe_match_indices(first_name: str, last_name: str, rsvp_people: List[Dict[str, str]]) -> List[int]:
    maybe_indices: List[int] = []

    for idx, person in enumerate(rsvp_people):
        direct_first_exact = bool(first_name and person["first"] and first_name == person["first"])
        direct_last_exact = bool(last_name and person["last"] and last_name == person["last"])
        swap_first_exact = bool(first_name and person["last"] and first_name == person["last"])
        swap_last_exact = bool(last_name and person["first"] and last_name == person["first"])

        direct_exact_full = direct_first_exact and direct_last_exact
        swap_exact_full = swap_first_exact and swap_last_exact
        if direct_exact_full or swap_exact_full:
            continue

        # Require one exact side and one similar side to qualify as "maybe".
        direct_first_similar = bool(first_name and person["first"] and is_similar_name(first_name, person["first"]))
        direct_last_similar = bool(last_name and person["last"] and is_similar_name(last_name, person["last"]))
        swap_first_similar = bool(first_name and person["last"] and is_similar_name(first_name, person["last"]))
        swap_last_similar = bool(last_name and person["first"] and is_similar_name(last_name, person["first"]))

        direct_maybe = (direct_first_exact and direct_last_similar) or (direct_last_exact and direct_first_similar)
        swap_maybe = (swap_first_exact and swap_last_similar) or (swap_last_exact and swap_first_similar)

        # Expanded fuzzy rule: allow maybe when both sides are similar but one part may
        # be embedded with extra tokens or minor typos in RSVP first/last representation.
        person_full = f"{person['first']} {person['last']}".strip()
        fuzzy_compound_maybe = (
            bool(last_name and is_similar_name(last_name, person["last"]) and has_similar_token(first_name, person_full))
            or bool(first_name and is_similar_name(first_name, person["first"]) and has_similar_token(last_name, person_full))
        )

        if direct_maybe or swap_maybe or fuzzy_compound_maybe:
            maybe_indices.append(idx)

    # Preserve order while removing duplicates.
    return list(dict.fromkeys(maybe_indices))


def find_exact_full_match_indices(first_name: str, last_name: str, rsvp_people: List[Dict[str, str]]) -> List[int]:
    exact_indices: List[int] = []
    if not first_name or not last_name:
        return exact_indices

    for idx, person in enumerate(rsvp_people):
        direct_exact = first_name == person["first"] and last_name == person["last"]
        swapped_exact = first_name == person["last"] and last_name == person["first"]
        if direct_exact or swapped_exact:
            exact_indices.append(idx)

    return list(dict.fromkeys(exact_indices))


def mark_matches(
    source_path: Path,
    output_path: Path,
    rsvp_people: List[Dict[str, str]],
    sheet_name: str = "Full List",
    status_col: int = 12,
    maybe_names_col: int = 13,
    review_col: int = 14,
    submission_count_col: int = 15,
) -> Tuple[int, int, int, int, int, int]:
    shutil.copy2(source_path, output_path)

    wb = load_workbook(output_path)
    if sheet_name not in wb.sheetnames:
        raise ValueError(
            f"Sheet '{sheet_name}' not found in {source_path.name}. "
            f"Available sheets: {', '.join(wb.sheetnames)}"
        )

    ws = wb[sheet_name]

    header_row = 1
    first_col, last_col, headers = detect_name_columns([cell.value for cell in ws[header_row]])
    if first_col is None or last_col is None:
        header_preview = ", ".join(f"{idx}:{name}" for idx, name in headers.items())
        raise ValueError(
            "Could not auto-detect first/last name columns on Full List sheet. "
            f"Detected headers: {header_preview}"
        )

    ws.cell(row=header_row, column=status_col).value = "Matched in RSVP"
    ws.cell(row=header_row, column=maybe_names_col).value = "Similar RSVP full names"
    ws.cell(row=header_row, column=review_col).value = "Unmatched RSVP names"
    ws.cell(row=header_row, column=submission_count_col).value = "Submission count"

    rsvp_first_names, rsvp_last_names = build_exact_name_sets(rsvp_people)

    total_rows = 0
    matched_rows = 0
    maybe_rows = 0
    first_hits = 0
    last_hits = 0
    matched_rsvp_indices: Set[int] = set()

    for row in range(header_row + 1, ws.max_row + 1):
        first_name = normalize_name(ws.cell(row=row, column=first_col).value)
        last_name = normalize_name(ws.cell(row=row, column=last_col).value)

        if not first_name and not last_name:
            continue

        total_rows += 1
        first_match = bool(first_name and first_name in rsvp_first_names)
        last_match = bool(last_name and last_name in rsvp_last_names)

        if first_match:
            first_hits += 1
        if last_match:
            last_hits += 1

        ws.cell(row=row, column=maybe_names_col).value = ""
        ws.cell(row=row, column=submission_count_col).value = ""
        exact_full_indices = find_exact_full_match_indices(first_name, last_name, rsvp_people)
        if exact_full_indices:
            ws.cell(row=row, column=status_col).value = "yes"
            ws.cell(row=row, column=submission_count_col).value = len(exact_full_indices)
            matched_rsvp_indices.update(exact_full_indices)
            matched_rows += 1
        else:
            maybe_indices = find_maybe_match_indices(first_name, last_name, rsvp_people)
            if maybe_indices:
                maybe_candidates = [rsvp_people[i]["full_display"] for i in maybe_indices]
                ws.cell(row=row, column=status_col).value = "maybe"
                ws.cell(row=row, column=maybe_names_col).value = "; ".join(maybe_candidates)
                ws.cell(row=row, column=submission_count_col).value = len(maybe_indices)
                matched_rsvp_indices.update(maybe_indices)
                maybe_rows += 1
            else:
                ws.cell(row=row, column=status_col).value = ""

    unmatched_rsvp_names = [
        person["full_display"] for idx, person in enumerate(rsvp_people) if idx not in matched_rsvp_indices
    ]
    unmatched_rsvp_names = list(dict.fromkeys(unmatched_rsvp_names))

    for row in range(header_row + 1, ws.max_row + 1):
        ws.cell(row=row, column=review_col).value = ""

    for offset, full_name in enumerate(unmatched_rsvp_names, start=0):
        ws.cell(row=header_row + 1 + offset, column=review_col).value = full_name

    wb.save(output_path)
    return total_rows, matched_rows, maybe_rows, first_hits, last_hits, len(unmatched_rsvp_names)


def parse_args() -> argparse.Namespace:
    script_dir = Path(__file__).resolve().parent

    parser = argparse.ArgumentParser(
        description=(
            "Copy AR7 attendee list and mark column L with 'yes' for exact matches or "
            "'maybe' for likely matches against RSVP columns G/H."
        )
    )
    parser.add_argument(
        "--source",
        default=None,
        help="Path to AR7 LAM3 Attendee List workbook. Defaults to pattern match in ./data.",
    )
    parser.add_argument(
        "--rsvp",
        default=None,
        help="Path to RSVP workbook. Defaults to pattern match in ./data.",
    )
    parser.add_argument(
        "--output",
        default=str(script_dir / "data" / "AR7 LAM3 Attendee List_marked.xlsx"),
        help="Output workbook path (copied from source and updated).",
    )
    parser.add_argument(
        "--sheet",
        default="Full List",
        help="Sheet name in source workbook to process.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    script_dir = Path(__file__).resolve().parent
    data_dir = script_dir / "data"

    source_path = resolve_input_path(data_dir, args.source, "AR7 LAM3 Attendee List.xlsx")
    rsvp_path = resolve_input_path(
        data_dir,
        args.rsvp,
        "IPCC AR7 WGII 3rd Lead Author Meeting RSVP*(1-209).xlsx",
    )
    output_path = Path(args.output)
    if not output_path.is_absolute():
        output_path = (script_dir / output_path).resolve()

    output_path.parent.mkdir(parents=True, exist_ok=True)
    rsvp_people = collect_rsvp_people(rsvp_path, first_col=7, last_col=8)

    total_rows, matched_rows, maybe_rows, first_hits, last_hits, unmatched_rsvp_count = mark_matches(
        source_path=source_path,
        output_path=output_path,
        rsvp_people=rsvp_people,
        sheet_name=args.sheet,
        status_col=12,
        maybe_names_col=13,
        review_col=14,
        submission_count_col=15,
    )

    print(f"Source: {source_path}")
    print(f"RSVP:   {rsvp_path}")
    print(f"Output: {output_path}")
    print(f"Processed rows: {total_rows}")
    print(f"Rows marked yes: {matched_rows}")
    print(f"Rows marked maybe: {maybe_rows}")
    print(f"First-name hits: {first_hits}")
    print(f"Last-name hits:  {last_hits}")
    print(f"Unmatched RSVP names: {unmatched_rsvp_count}")


if __name__ == "__main__":
    main()
